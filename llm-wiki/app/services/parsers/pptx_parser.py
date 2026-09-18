# -*- coding: utf-8 -*-
"""pptx 解析器：python-pptx 遍历幻灯片文本框 → 纯文本"""
from __future__ import annotations

import os
from io import BytesIO

from pptx import Presentation

from .base import DocumentParser, ParsedDocument, register_parser


@register_parser
class PptxParser(DocumentParser):
    extensions = {".pptx"}

    def parse_text(self, raw: str, file_name: str = "") -> ParsedDocument:
        raise NotImplementedError("pptx 为二进制格式，请使用 parse_bytes")

    def parse_bytes(self, data: bytes, file_name: str = "") -> ParsedDocument:
        try:
            prs = Presentation(BytesIO(data))
        except Exception as e:  # noqa: BLE001
            raise ValueError(f"pptx 解析失败: {e}") from e

        slide_texts: list[str] = []
        for slide in prs.slides:
            shape_texts: list[str] = []
            for shape in slide.shapes:
                if not getattr(shape, "has_text_frame", False):
                    continue
                text = (shape.text_frame.text or "").strip()
                if text:
                    shape_texts.append(text)
            if shape_texts:
                slide_texts.append("\n".join(shape_texts))

        content = "\n\n".join(slide_texts)
        title = os.path.splitext(file_name)[0]
        return ParsedDocument(
            title=title, metadata={"type": "doc", "tags": []},
            content=content, raw="",
        )

    def serialize(self, doc: ParsedDocument) -> str:
        raise NotImplementedError("pptx 解析器不支持反向序列化")
