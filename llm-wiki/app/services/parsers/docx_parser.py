# -*- coding: utf-8 -*-
"""docx 解析器：python-docx 读取段落 + 表格 → 纯文本

只支持二进制解析（.docx 本质为 zip+xml），由 python-docx 内部完成解码，
此处不直接处理编码问题。
"""
from __future__ import annotations

import os
from io import BytesIO

from docx import Document

from .base import DocumentParser, ParsedDocument, register_parser


@register_parser
class DocxParser(DocumentParser):
    extensions = {".docx"}

    def parse_text(self, raw: str, file_name: str = "") -> ParsedDocument:
        raise NotImplementedError("docx 为二进制格式，请使用 parse_bytes")

    def parse_bytes(self, data: bytes, file_name: str = "") -> ParsedDocument:
        """段落逐行输出，空段落留作换行分隔；表格每行用竖线连接"""
        try:
            document = Document(BytesIO(data))
        except Exception as e:  # noqa: BLE001
            raise ValueError(f"docx 解析失败: {e}") from e

        lines: list[str] = []
        for para in document.paragraphs:
            # 空段落保留为空行，用于段落间换行分隔
            lines.append((para.text or "").strip())

        for table in document.tables:
            for row in table.rows:
                cells = [(cell.text or "").strip() for cell in row.cells]
                lines.append(" | ".join(cells))

        content = "\n".join(lines).strip()
        title = os.path.splitext(file_name)[0]
        return ParsedDocument(
            title=title, metadata={"type": "doc", "tags": []},
            content=content, raw="",
        )

    def serialize(self, doc: ParsedDocument) -> str:
        raise NotImplementedError("docx 解析器不支持反向序列化")
