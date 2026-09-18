# -*- coding: utf-8 -*-
"""xlsx 解析器：openpyxl 只读遍历工作表 → 纯文本

单元格值以 tab 连接为一行；空行跳过；sheet 之间加分隔标题。
"""
from __future__ import annotations

import os
from io import BytesIO

from openpyxl import load_workbook

from .base import DocumentParser, ParsedDocument, register_parser


@register_parser
class XlsxParser(DocumentParser):
    extensions = {".xlsx"}

    def parse_text(self, raw: str, file_name: str = "") -> ParsedDocument:
        raise NotImplementedError("xlsx 为二进制格式，请使用 parse_bytes")

    def parse_bytes(self, data: bytes, file_name: str = "") -> ParsedDocument:
        try:
            wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            raise ValueError(f"xlsx 解析失败: {e}") from e

        blocks: list[str] = []
        for ws in wb.worksheets:
            lines = [f"## {ws.title}"]
            for row in ws.iter_rows(values_only=True):
                cells = [str(c).strip() for c in row if c is not None]
                if cells:
                    lines.append("\t".join(cells))
            blocks.append("\n".join(lines))
        wb.close()

        content = "\n\n".join(blocks)
        title = os.path.splitext(file_name)[0]
        return ParsedDocument(
            title=title, metadata={"type": "doc", "tags": []},
            content=content, raw="",
        )

    def serialize(self, doc: ParsedDocument) -> str:
        raise NotImplementedError("xlsx 解析器不支持反向序列化")
